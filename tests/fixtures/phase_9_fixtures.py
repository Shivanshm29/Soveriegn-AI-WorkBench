"""Test fixtures and sample assets for Phase 9 testing and demonstration scenarios."""

import os
import openpyxl
from PIL import Image, ImageDraw
from typing import Dict, Any, List

SAMPLE_CSV_CONTENT = """equipment_id,timestamp,vibration_mms,temperature_c,pressure_bar,status
PX-417,2026-09-01T08:00:00Z,2.1,65.4,5.8,NORMAL
PX-417,2026-09-01T12:00:00Z,2.3,67.1,5.9,NORMAL
PX-417,2026-09-01T16:00:00Z,4.8,78.5,6.4,ELEVATED
PX-417,2026-09-01T20:00:00Z,5.2,82.0,6.7,ALERT
PX-418,2026-09-01T08:00:00Z,1.8,62.0,5.5,NORMAL
PX-418,2026-09-01T12:00:00Z,1.9,63.2,5.6,NORMAL
"""

SAMPLE_CODE_NORMAL = """# Calculate sum and mean of numbers
data = [12.5, 14.2, 11.8, 15.0, 13.5]
total = sum(data)
avg = total / len(data)
print(f"COUNT: {len(data)}")
print(f"SUM: {total:.2f}")
print(f"MEAN: {avg:.2f}")
"""

SAMPLE_CODE_NETWORK_ATTEMPT = """import urllib.request
try:
    urllib.request.urlopen("https://example.com", timeout=3)
    print("NETWORK_SUCCESS")
except Exception as e:
    print(f"NETWORK_BLOCKED_CAUGHT: {e}")
"""

SAMPLE_CODE_PATH_TRAVERSAL = """import os
try:
    with open("../../secret.txt", "r") as f:
        print("SECRET_LEAKED:", f.read())
except Exception as e:
    print(f"ACCESS_ERROR: {e}")
"""

SAMPLE_CODE_ENV_ACCESS = """import os
try:
    with open(".env", "r") as f:
        print("ENV_LEAKED:", f.read())
except Exception as e:
    print(f"ENV_ERROR: {e}")
"""

SAMPLE_CODE_TIMEOUT = """import time
time.sleep(100)
"""

SAMPLE_FINDINGS = [
    {
        "observation": "Candidate surface pitting indication observed on flange rim",
        "region": "Flange body / Quadrant 2",
        "confidence": "MEDIUM",
        "evidence_id": "EV-VIS-001",
    },
    {
        "observation": "Nominal diameter notation indicated as 250mm +/- 0.5mm",
        "region": "Drawing Dimension Block",
        "confidence": "HIGH",
        "evidence_id": "EV-DIM-002",
    },
]

SAMPLE_EVIDENCE_CITATIONS = [
    {
        "chunk_id": "doc_pump_man_417_c1",
        "source_document": "PX417_Maintenance_SOP.pdf",
        "source_hash": "a1b2c3d4e5f67890123456789abcdef0123456789abcdef0123456789abcdef0",
        "page": 4,
    },
    {
        "chunk_id": "doc_insp_flange_c2",
        "source_document": "Turbine_Flange_Inspection_Scan.pdf",
        "source_hash": "b2c3d4e5f67890123456789abcdef0123456789abcdef0123456789abcdef01",
        "page": 1,
    },
]


def create_sample_csv_file(path: str) -> str:
    """Write sample industrial measurement CSV to specified path."""
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(SAMPLE_CSV_CONTENT)
    return path


def create_sample_xlsx_file(path: str) -> str:
    """Create sample industrial measurement Excel workbook."""
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Measurements"

    lines = [line.strip().split(",") for line in SAMPLE_CSV_CONTENT.strip().splitlines()]
    for r_idx, row in enumerate(lines, start=1):
        for c_idx, val in enumerate(row, start=1):
            try:
                num = float(val)
                ws.cell(row=r_idx, column=c_idx, value=num)
            except ValueError:
                ws.cell(row=r_idx, column=c_idx, value=val)

    wb.save(path)
    wb.close()
    return path


def create_industrial_equipment_photo(path: str, width: int = 800, height: int = 600) -> str:
    """Create a sample industrial equipment inspection photo / diagram."""
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    img = Image.new("RGB", (width, height), color=(240, 243, 246))
    draw = ImageDraw.Draw(img)

    # Frame and Header
    draw.rectangle([20, 20, width - 20, height - 20], outline=(40, 50, 60), width=3)
    draw.rectangle([20, 20, width - 20, 80], fill=(30, 41, 59))
    draw.text((40, 38), "INDUSTRIAL INSPECTION REPORT: TURBINE-01 / CASING", fill=(255, 255, 255))

    # Casing / Flange geometry
    draw.rectangle([100, 140, width - 100, height - 120], outline=(70, 80, 95), width=4, fill=(215, 225, 235))
    draw.ellipse([250, 200, 550, 480], outline=(30, 40, 50), width=3, fill=(180, 195, 210))
    draw.ellipse([340, 280, 460, 400], outline=(15, 23, 42), width=2, fill=(130, 145, 160))

    # Inspection anomaly indicator
    draw.rectangle([430, 230, 520, 310], outline=(220, 38, 38), width=3)
    draw.text((435, 210), "[DEFECT-IND-01]", fill=(185, 28, 28))
    draw.text((435, 320), "Candidate surface pitting", fill=(185, 28, 28))

    # Measurement Callouts
    draw.text((50, 100), "Nominal Diameter: 250mm +/- 0.5mm", fill=(30, 40, 50))
    draw.text((50, height - 80), "STATUS: NON-CONFORMANCE CANDIDATE DETECTED | NDT REQUIRED", fill=(185, 28, 28))

    img.save(path)
    return path

