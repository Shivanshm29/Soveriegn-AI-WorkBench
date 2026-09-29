"""Local XLSX workbook generator creating styled spreadsheets with calculations and provenance."""

import os
import hashlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, List, Optional, Any

from backend.app.artifacts.schemas import (
    ArtifactMetadata,
    SpreadsheetContent,
    SpreadsheetSheetData,
)


class XlsxArtifactGenerator:
    """Generates styled Excel workbooks locally via openpyxl."""

    @staticmethod
    def _compute_file_hash(path: str) -> str:
        """Compute SHA-256 content hash of file."""
        hasher = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        return hasher.hexdigest()

    def generate_workbook(
        self,
        content: SpreadsheetContent,
        output_path: str,
    ) -> ArtifactMetadata:
        """Create multi-sheet Excel workbook with styled tables, formulas, and provenance metadata."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for sheet_data in content.sheets:
            ws = wb.create_sheet(title=sheet_data.title[:30])

            # Title
            ws.cell(row=1, column=1, value=sheet_data.title).font = title_font
            ws.row_dimensions[1].height = 24

            # Headers
            start_row = 3
            for col_idx, header in enumerate(sheet_data.headers, start=1):
                cell = ws.cell(row=start_row, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws.row_dimensions[start_row].height = 20

            # Data rows
            current_row = start_row + 1
            for row_vals in sheet_data.rows:
                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.border = thin_border
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")
                current_row += 1

            # Summary formula row if requested
            if sheet_data.include_summary_row and sheet_data.rows:
                ws.cell(row=current_row, column=1, value="AVERAGE / SUMMARY").font = Font(bold=True)
                for c_idx in range(2, len(sheet_data.headers) + 1):
                    col_let = get_column_letter(c_idx)
                    # Check if numeric column
                    first_val = sheet_data.rows[0][c_idx - 1] if len(sheet_data.rows[0]) >= c_idx else None
                    if isinstance(first_val, (int, float)):
                        formula = f"=AVERAGE({col_let}{start_row + 1}:{col_let}{current_row - 1})"
                        sum_cell = ws.cell(row=current_row, column=c_idx, value=formula)
                        sum_cell.font = Font(bold=True)
                        sum_cell.border = thin_border
                current_row += 1

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Provenance Sheet
        prov_ws = wb.create_sheet(title="Provenance")
        prov_ws.cell(row=1, column=1, value="Sovereign Audit & Provenance").font = title_font
        prov_ws.cell(row=3, column=1, value="Task ID:").font = Font(bold=True)
        prov_ws.cell(row=3, column=2, value=content.task_id)
        prov_ws.cell(row=4, column=1, value="Workbook Title:").font = Font(bold=True)
        prov_ws.cell(row=4, column=2, value=content.title)

        p_row = 6
        for k, v in content.provenance.items():
            prov_ws.cell(row=p_row, column=1, value=str(k)).font = Font(bold=True)
            prov_ws.cell(row=p_row, column=2, value=str(v))
            p_row += 1

        wb.save(output_path)

        content_hash = self._compute_file_hash(output_path)
        file_size = os.path.getsize(output_path)

        return ArtifactMetadata(
            task_id=content.task_id,
            artifact_type="xlsx",
            file_path=output_path,
            content_hash=content_hash,
            file_size_bytes=file_size,
            status="GENERATED",
            metadata={"title": content.title, "sheets_count": len(content.sheets)},
        )
