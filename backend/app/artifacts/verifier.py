"""Deterministic verifier for generated artifacts (DOCX, XLSX, TXT, JSON)."""

import os
import hashlib
import docx
import openpyxl
from typing import Dict, List, Optional, Any

from backend.app.artifacts.schemas import (
    ArtifactMetadata,
    ArtifactVerificationResult,
)


class ArtifactVerifier:
    """Performs deterministic verification of generated artifacts."""

    @staticmethod
    def _compute_hash(path: str) -> str:
        hasher = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        return hasher.hexdigest()

    @classmethod
    def verify(
        cls,
        artifact: ArtifactMetadata,
        known_evidence_hashes: Optional[List[str]] = None,
    ) -> ArtifactVerificationResult:
        """Verify artifact integrity, readable format, valid hash, and non-fabricated evidence."""
        checks: Dict[str, bool] = {}
        failure_reasons: List[str] = []
        details: Dict[str, Any] = {}

        # 1. File existence
        exists = os.path.exists(artifact.file_path)
        checks["file_exists"] = exists
        if not exists:
            failure_reasons.append(f"Artifact file not found at: {artifact.file_path}")
            return ArtifactVerificationResult(
                is_verified=False,
                artifact_id=artifact.artifact_id,
                artifact_type=artifact.artifact_type,
                checks=checks,
                failure_reasons=failure_reasons,
                details=details,
            )

        # 2. Non-zero size
        size = os.path.getsize(artifact.file_path)
        checks["non_zero_size"] = size > 0
        details["file_size_bytes"] = size
        if size == 0:
            failure_reasons.append("Artifact file size is 0 bytes.")

        # 3. Content hash integrity
        computed_hash = cls._compute_hash(artifact.file_path)
        hash_matches = (computed_hash == artifact.content_hash)
        checks["hash_matches"] = hash_matches
        details["computed_hash"] = computed_hash
        details["expected_hash"] = artifact.content_hash
        if not hash_matches:
            failure_reasons.append(f"Content hash mismatch: expected {artifact.content_hash}, got {computed_hash}")

        # 4. Task ID presence
        task_id_valid = bool(artifact.task_id and artifact.task_id.strip())
        checks["task_id_valid"] = task_id_valid
        if not task_id_valid:
            failure_reasons.append("Artifact missing valid task_id.")

        # 5. Format reopening verification
        if artifact.artifact_type == "docx":
            try:
                doc = docx.Document(artifact.file_path)
                paragraph_count = len(doc.paragraphs)
                table_count = len(doc.tables)
                checks["docx_reopened_ok"] = True
                details["paragraph_count"] = paragraph_count
                details["table_count"] = table_count
            except Exception as e:
                checks["docx_reopened_ok"] = False
                failure_reasons.append(f"Failed to reopen DOCX artifact: {e}")

        elif artifact.artifact_type == "xlsx":
            try:
                wb = openpyxl.load_workbook(artifact.file_path, data_only=True)
                sheet_count = len(wb.sheetnames)
                wb.close()
                checks["xlsx_reopened_ok"] = True
                details["sheet_count"] = sheet_count
            except Exception as e:
                checks["xlsx_reopened_ok"] = False
                failure_reasons.append(f"Failed to reopen XLSX artifact: {e}")

        # 6. Evidence citations check (no fabricated IDs)
        if artifact.source_evidence:
            invalid_citations = []
            for ev in artifact.source_evidence:
                ev_id = ev.get("chunk_id") or ev.get("evidence_id")
                if not ev_id or ev_id.strip() == "":
                    invalid_citations.append(ev)
                elif known_evidence_hashes:
                    s_hash = ev.get("source_hash") or ev.get("sha256")
                    if s_hash and s_hash not in known_evidence_hashes:
                        invalid_citations.append(ev)

            checks["evidence_citations_valid"] = (len(invalid_citations) == 0)
            if invalid_citations:
                failure_reasons.append(f"Detected invalid or fabricated evidence citations: {invalid_citations}")
        else:
            checks["evidence_citations_valid"] = True

        is_verified = (len(failure_reasons) == 0) and all(checks.values())
        return ArtifactVerificationResult(
            is_verified=is_verified,
            artifact_id=artifact.artifact_id,
            artifact_type=artifact.artifact_type,
            checks=checks,
            failure_reasons=failure_reasons,
            details=details,
        )
