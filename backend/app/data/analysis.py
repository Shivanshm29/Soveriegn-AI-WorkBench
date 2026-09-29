"""Local tabular dataset analyzer supporting CSV and XLSX files with safety limits."""

import csv
import io
import os
import openpyxl
from typing import Dict, List, Optional, Any, Tuple, Union

from backend.app.data.schemas import (
    DataAnalysisRequest,
    DataAnalysisResult,
    ColumnStats,
    TableData,
    CalculationTraceStep,
)


class DataAnalysisError(Exception):
    """Raised when data reading or analysis operations fail."""
    pass


class DataAnalyzer:
    """Performs deterministic analysis, summaries, filtering, and aggregation on CSV and XLSX datasets."""

    def __init__(self, max_rows: int = 50_000, max_file_size_bytes: int = 25_000_000):
        self.max_rows = max_rows
        self.max_file_size_bytes = max_file_size_bytes

    def read_csv(self, file_path_or_content: Union[str, bytes], is_raw: bool = False) -> TableData:
        """Read CSV data into normalized TableData format."""
        if is_raw:
            text = file_path_or_content.decode("utf-8") if isinstance(file_path_or_content, bytes) else file_path_or_content
            f = io.StringIO(text)
        else:
            path = str(file_path_or_content)
            if not os.path.exists(path):
                raise FileNotFoundError(f"CSV file not found: {path}")
            if os.path.getsize(path) > self.max_file_size_bytes:
                raise DataAnalysisError(f"File size exceeds maximum limit of {self.max_file_size_bytes} bytes.")
            f = open(path, "r", encoding="utf-8", errors="replace")

        try:
            reader = csv.reader(f)
            headers = next(reader, [])
            rows = []
            for idx, row in enumerate(reader):
                if idx >= self.max_rows:
                    break
                rows.append(row)
            return TableData(
                headers=headers,
                rows=rows,
                row_count=len(rows),
                column_count=len(headers),
            )
        finally:
            if not is_raw and hasattr(f, "close"):
                f.close()

    def read_xlsx(self, file_path: str, sheet_name: Optional[str] = None) -> TableData:
        """Read XLSX workbook into normalized TableData format."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"XLSX file not found: {file_path}")
        if os.path.getsize(file_path) > self.max_file_size_bytes:
            raise DataAnalysisError(f"File size exceeds maximum limit of {self.max_file_size_bytes} bytes.")

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        headers = []
        rows = []
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if r_idx == 0:
                headers = [str(cell) if cell is not None else f"col_{i}" for i, cell in enumerate(row)]
            else:
                if len(rows) >= self.max_rows:
                    break
                rows.append(list(row))

        wb.close()
        return TableData(
            headers=headers,
            rows=rows,
            row_count=len(rows),
            column_count=len(headers),
        )

    def summarize(self, data: TableData) -> Dict[str, ColumnStats]:
        """Compute statistical summary for each column in TableData."""
        stats: Dict[str, ColumnStats] = {}
        for c_idx, h in enumerate(data.headers):
            vals = []
            null_count = 0
            for r in data.rows:
                if c_idx < len(r):
                    val = r[c_idx]
                    if val is None or val == "" or str(val).lower() in ("nan", "null", "none"):
                        null_count += 1
                    else:
                        try:
                            vals.append(float(val))
                        except (ValueError, TypeError):
                            vals.append(str(val))
                else:
                    null_count += 1

            num_vals = [v for v in vals if isinstance(v, (int, float))]
            if len(num_vals) > 0 and len(num_vals) >= len(vals) * 0.8:
                dtype = "numeric"
                min_v = min(num_vals)
                max_v = max(num_vals)
                sum_v = sum(num_vals)
                mean_v = sum_v / len(num_vals)
            else:
                dtype = "string"
                min_v = None
                max_v = None
                sum_v = None
                mean_v = None

            stats[h] = ColumnStats(
                name=h,
                dtype=dtype,
                total_count=data.row_count,
                null_count=null_count,
                min_value=round(min_v, 4) if min_v is not None else None,
                max_value=round(max_v, 4) if max_v is not None else None,
                mean_value=round(mean_v, 4) if mean_v is not None else None,
                sum_value=round(sum_v, 4) if sum_v is not None else None,
            )
        return stats

    def filter(
        self,
        data: TableData,
        column: str,
        operator: str,
        target_value: Any,
    ) -> TableData:
        """Filter TableData rows based on column condition."""
        if column not in data.headers:
            raise DataAnalysisError(f"Filter column '{column}' not found in headers {data.headers}")
        c_idx = data.headers.index(column)

        filtered_rows = []
        for r in data.rows:
            if c_idx >= len(r):
                continue
            val = r[c_idx]
            match = False
            try:
                # Numerical comparison if possible
                num_v = float(val)
                num_target = float(target_value)
                if operator == ">":
                    match = num_v > num_target
                elif operator == ">=":
                    match = num_v >= num_target
                elif operator == "<":
                    match = num_v < num_target
                elif operator == "<=":
                    match = num_v <= num_target
                elif operator in ("==", "="):
                    match = num_v == num_target
                elif operator in ("!=", "<>"):
                    match = num_v != num_target
            except (ValueError, TypeError):
                # String comparison
                s_v = str(val).strip().lower()
                s_target = str(target_value).strip().lower()
                if operator in ("==", "="):
                    match = s_v == s_target
                elif operator in ("!=", "<>"):
                    match = s_v != s_target
                elif operator == "contains":
                    match = s_target in s_v

            if match:
                filtered_rows.append(r)

        return TableData(
            headers=data.headers,
            rows=filtered_rows,
            row_count=len(filtered_rows),
            column_count=len(data.headers),
        )

    def group_and_aggregate(
        self,
        data: TableData,
        group_column: str,
        aggregate_column: str,
        func: str = "mean",
    ) -> Dict[str, float]:
        """Group rows by a key column and compute aggregation on target column."""
        if group_column not in data.headers:
            raise DataAnalysisError(f"Group column '{group_column}' not found.")
        if aggregate_column not in data.headers:
            raise DataAnalysisError(f"Aggregate column '{aggregate_column}' not found.")

        g_idx = data.headers.index(group_column)
        a_idx = data.headers.index(aggregate_column)

        groups: Dict[str, List[float]] = {}
        for r in data.rows:
            if g_idx < len(r) and a_idx < len(r):
                g_key = str(r[g_idx]).strip()
                try:
                    num_val = float(r[a_idx])
                    groups.setdefault(g_key, []).append(num_val)
                except (ValueError, TypeError):
                    continue

        results: Dict[str, float] = {}
        for g_key, vals in groups.items():
            if not vals:
                continue
            if func.lower() in ("mean", "avg", "average"):
                results[g_key] = round(sum(vals) / len(vals), 4)
            elif func.lower() == "sum":
                results[g_key] = round(sum(vals), 4)
            elif func.lower() == "count":
                results[g_key] = len(vals)
            elif func.lower() == "min":
                results[g_key] = min(vals)
            elif func.lower() == "max":
                results[g_key] = max(vals)

        return results

    def analyze(self, request: DataAnalysisRequest) -> DataAnalysisResult:
        """Execute end-to-end dataset analysis according to request specifications."""
        traces: List[CalculationTraceStep] = []

        try:
            if request.file_path and request.file_path.lower().endswith(".xlsx"):
                data = self.read_xlsx(request.file_path)
                f_type = "xlsx"
            elif request.file_path:
                data = self.read_csv(request.file_path)
                f_type = "csv"
            elif request.raw_content:
                data = self.read_csv(request.raw_content, is_raw=True)
                f_type = "csv"
            else:
                raise DataAnalysisError("Neither file_path nor raw_content provided.")

            column_stats = self.summarize(data)
            traces.append(
                CalculationTraceStep(
                    step="summarize_dataset",
                    formula="compute_column_stats(dataset)",
                    inputs={"rows": data.row_count, "columns": len(data.headers)},
                    output={"columns": data.headers},
                )
            )

            aggregated_results: Dict[str, Any] = {}
            if request.group_by_column and request.aggregate_column:
                agg_func = request.aggregate_function or "mean"
                aggregated_results = self.group_and_aggregate(
                    data,
                    request.group_by_column,
                    request.aggregate_column,
                    agg_func,
                )
                traces.append(
                    CalculationTraceStep(
                        step="group_and_aggregate",
                        formula=f"{agg_func}({request.aggregate_column}) GROUP BY {request.group_by_column}",
                        inputs={"group": request.group_by_column, "target": request.aggregate_column},
                        output=aggregated_results,
                    )
                )

            filtered_count = None
            filtered_sample = None
            if request.filter_column and request.filter_value is not None:
                filtered_data = self.filter(
                    data,
                    request.filter_column,
                    request.filter_operator or "==",
                    request.filter_value,
                )
                filtered_count = filtered_data.row_count
                filtered_sample = [
                    dict(zip(filtered_data.headers, r)) for r in filtered_data.rows[:5]
                ]
                traces.append(
                    CalculationTraceStep(
                        step="filter_rows",
                        formula=f"{request.filter_column} {request.filter_operator} {request.filter_value}",
                        inputs={"initial_rows": data.row_count},
                        output={"retained_rows": filtered_count},
                    )
                )

            return DataAnalysisResult(
                task_id=request.task_id,
                file_type=f_type,
                row_count=data.row_count,
                column_count=data.column_count,
                columns=data.headers,
                column_stats=column_stats,
                aggregated_results=aggregated_results,
                filtered_row_count=filtered_count,
                filtered_rows_sample=filtered_sample,
                calculation_traces=traces,
                status="SUCCESS",
            )
        except Exception as e:
            return DataAnalysisResult(
                task_id=request.task_id,
                file_type=request.file_type,
                row_count=0,
                column_count=0,
                columns=[],
                status="FAILED",
                error=str(e),
            )
